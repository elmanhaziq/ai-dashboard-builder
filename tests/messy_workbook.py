"""An extremely messy Excel workbook, the way real business reports look. Used by tests and the sample-data script.

Sheets:
* **Sales Report** — company title and a note above the data; a sales table with a merged two-row header, mixed date
  formats, inconsistent region labels, a stray blank row and a Total row; a second "Targets" table beside it; a
  footnote underneath.
* **Monthly** — a "wide" report: one column per month plus a Total column and a Total row.
* **Branch Profile** — a sideways table: field names down column A, one branch per column.
* **Clean** — an ordinary table (should be loaded exactly as before).
"""
from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook

REGIONS = ["Central", "central", "CENTRAL ", "Northern", "Southern", "East Coast", "east coast", "Sabah"]


def build() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales Report"
    ws["A1"] = "ACME Malaysia Sdn Bhd — Sales Report 2024"
    ws["A2"] = "Prepared by Finance, confidential"
    # two-row header with a merged group label
    ws["D5"] = "Amount"
    ws.merge_cells("D5:E5")
    for j, h in enumerate(["Order ID", "Order Date", "Region", "Sales (RM)", "Units"], start=1):
        ws.cell(row=6, column=j, value=h)
    r = 7
    for i in range(30):
        if i == 15:
            r += 1  # a stray blank row inside the table
        d = date(2024, 1 + i % 6, 1 + (i * 3) % 27)
        date_value = d if i % 3 == 0 else (d.strftime("%d/%m/%Y") if i % 3 == 1 else d.strftime("%Y-%m-%d"))
        sales = f"RM {1000 + i * 37:,}.50" if i % 4 == 0 else 1000 + i * 37 + 0.5
        ws.cell(row=r, column=1, value=5001 + i)
        ws.cell(row=r, column=2, value=date_value)
        ws.cell(row=r, column=3, value=REGIONS[i % len(REGIONS)])
        ws.cell(row=r, column=4, value=None if i == 7 else sales)
        ws.cell(row=r, column=5, value=1 + i % 9)
        r += 1
    ws.cell(row=r, column=3, value="Total")
    ws.cell(row=r, column=4, value=99999)
    ws.cell(row=r + 3, column=1, value="Note: figures exclude SST. Source: ERP export.")
    # a second table beside the first, with its own title
    ws["H4"] = "Targets"
    ws["H5"], ws["I5"] = "Region", "Target (RM)"
    for i, (reg, tgt) in enumerate([("Central", 12000), ("Northern", 8000), ("Southern", 9000), ("East Coast", 7000)]):
        ws.cell(row=6 + i, column=8, value=reg)
        ws.cell(row=6 + i, column=9, value=tgt)

    ws2 = wb.create_sheet("Monthly")
    ws2["A1"] = "Monthly sales by region (RM) 2024"
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun"]
    for j, h in enumerate(["Region"] + months + ["Total"], start=1):
        ws2.cell(row=3, column=j, value=h)
    for i, reg in enumerate(["Central", "Northern", "Southern", "East Coast", "Sabah"]):
        vals = [1000 + 100 * i + 50 * m for m in range(6)]
        ws2.cell(row=4 + i, column=1, value=reg)
        for m, v in enumerate(vals):
            ws2.cell(row=4 + i, column=2 + m, value=v)
        ws2.cell(row=4 + i, column=8, value=sum(vals))
    ws2.cell(row=9, column=1, value="Total")
    for m in range(7):
        ws2.cell(row=9, column=2 + m, value=0)

    ws3 = wb.create_sheet("Branch Profile")
    branches = ["KL Central", "Penang", "Johor Bahru", "Kuching"]
    for j, b in enumerate(branches, start=2):
        ws3.cell(row=1, column=j, value=b)
    fields = [("Branch code", ["KL01", "PG03", "JB04", "KC07"]),
              ("Opened", [date(2015, 3, 1), date(2017, 6, 15), date(2019, 1, 10), date(2021, 9, 1)]),
              ("Staff", [42, 25, 31, 12]),
              ("Manager", ["Aisyah", "Ravi", "Mei Ling", "Daniel"]),
              ("Monthly rent (RM)", [18500, 9200, 11000, 6400])]
    for i, (name, values) in enumerate(fields, start=2):
        ws3.cell(row=i, column=1, value=name)
        for j, v in enumerate(values, start=2):
            ws3.cell(row=i, column=j, value=v)

    ws4 = wb.create_sheet("Clean")
    ws4.append(["Date", "Product", "Qty", "Price"])
    for i in range(20):
        ws4.append([date(2024, 1, 1 + i), f"P{i % 4}", i + 1, 9.9 + i])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
